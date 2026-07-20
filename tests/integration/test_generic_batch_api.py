import io
import json
import zipfile

import openpyxl
from fastapi.testclient import TestClient

from datawork.web.app import create_app


CSV = (
    "批次标签,任意分组,响应变量\n"
    "第一批,A,1\n第一批,A,2\n第一批,A,3\n"
    "第一批,B,4\n第一批,B,5\n第一批,B,6\n"
    "第二批,A,2\n第二批,A,3\n第二批,A,4\n"
    "第二批,B,5\n第二批,B,6\n第二批,B,7\n"
).encode("utf-8")


def _plan():
    return {
        "dependent_variables": ["响应变量"],
        "fixed_factors": ["任意分组"],
        "split_by": ["批次标签"],
        "method": "welch_ttest",
        "alpha": 0.05,
        "ss_type": 3,
    }


def test_preflight_blocks_missing_selections(tmp_path):
    client = TestClient(create_app(workspace_root=tmp_path / "workspace"))
    response = client.post(
        "/api/preflight",
        files={"file": ("任意数据.csv", CSV, "text/csv")},
        data={"plan_json": json.dumps({"method": "welch_ttest"}, ensure_ascii=False)},
    )
    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["ready"] is False
    fields = {item["field"] for item in payload["issues"]}
    assert "dependent_variables" in fields
    assert "fixed_factors" in fields


def test_generic_batch_uses_user_selected_columns_and_exports_multi_sheet_workbook(tmp_path):
    client = TestClient(create_app(workspace_root=tmp_path / "workspace"))

    profile = client.post(
        "/api/profile",
        files={"file": ("任意数据.csv", CSV, "text/csv")},
    )
    assert profile.status_code == 200, profile.text
    profile_payload = profile.json()
    assert "batch_workflow" not in profile_payload
    assert profile_payload["columns"] == ["批次标签", "任意分组", "响应变量"]

    preflight = client.post(
        "/api/preflight",
        files={"file": ("任意数据.csv", CSV, "text/csv")},
        data={"plan_json": json.dumps(_plan(), ensure_ascii=False)},
    )
    assert preflight.status_code == 200, preflight.text
    check = preflight.json()
    assert check["ready"] is True
    assert check["batch_summary"]["group_count"] == 2
    assert check["batch_summary"]["split_by"] == ["批次标签"]

    response = client.post(
        "/api/analyze",
        files={"file": ("任意数据.csv", CSV, "text/csv")},
        data={"plan_json": json.dumps(_plan(), ensure_ascii=False)},
    )
    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["kind"] == "batch"
    assert payload["result"]["split_cols"] == ["批次标签"]
    assert len(payload["result"]["summary"]) == 2
    assert payload["batch_export"]["status"] == "preparing"
    assert payload["batch_export"]["status_url"].endswith("/status")
    assert payload["batch_export"]["sheet_name"] == "结果总览"
    assert payload["batch_export"]["sheet_names"] == ["结果总览", "批次_001", "批次_002", "失败与警告", "分析设置"]

    export_status = client.get(payload["batch_export"]["status_url"])
    assert export_status.status_code == 200, export_status.text
    assert export_status.json()["status"] == "ready"
    assert export_status.json()["standardized_report"]["artifacts"]

    download = client.get(payload["batch_export"]["download_url"])
    assert download.status_code == 200
    workbook = openpyxl.load_workbook(io.BytesIO(download.content), read_only=True)
    assert workbook.sheetnames == ["结果总览", "批次_001", "批次_002", "失败与警告", "分析设置"]
    worksheet = workbook["结果总览"]
    headers = [cell.value for cell in next(worksheet.iter_rows(min_row=3, max_row=3))]
    assert "批次标签" in headers
    assert "样本量" in headers


def test_factor_combination_api_generates_unique_models_and_batch_sheets(tmp_path):
    rows = ["A,B,C,D,y"]
    for a in ["0", "1"]:
        for b in ["0", "1"]:
            for c in ["0", "1"]:
                for d in ["0", "1"]:
                    for repeat in range(4):
                        value = int(a) + 2 * int(b) + 3 * int(c) + 4 * int(d) + repeat * 0.1
                        rows.append(f"{a},{b},{c},{d},{value}")
    csv_bytes = ("\n".join(rows) + "\n").encode()
    plan = {
        "dependent_variables": ["y"],
        "fixed_factors": ["A", "B", "C", "D"],
        "method": "twoway_anova",
        "factor_combinations_enabled": True,
        "factor_combination_min_order": 2,
        "factor_combination_max_order": 2,
        "combination_p_adjust": "holm",
    }
    client = TestClient(create_app(workspace_root=tmp_path / "workspace"))
    preflight = client.post(
        "/api/preflight",
        files={"file": ("combinations.csv", csv_bytes, "text/csv")},
        data={"plan_json": json.dumps(plan)},
    )
    assert preflight.status_code == 200, preflight.text
    check = preflight.json()
    assert check["ready"] is True
    assert check["batch_summary"]["factor_combination_count"] == 6

    response = client.post(
        "/api/analyze",
        files={"file": ("combinations.csv", csv_bytes, "text/csv")},
        data={"plan_json": json.dumps(plan)},
    )
    assert response.status_code == 200, response.text
    payload = response.json()
    combinations = {row["factor_combination"] for row in payload["result"]["summary"]}
    assert combinations == {"A × B", "A × C", "A × D", "B × C", "B × D", "C × D"}
    assert payload["batch_export"]["sheet_name"] == "结果总览"
    assert payload["batch_export"]["sheet_names"][:3] == ["结果总览", "批次_001", "批次_002"]
    assert all(row["p_adjust_method"] == "holm" for row in payload["result"]["summary"])


def test_none_adjustment_and_custom_combination_names_flow_through_api_and_export(tmp_path):
    rows = ["A,B,y"]
    for a in ["0", "1"]:
        for b in ["0", "1"]:
            for repeat in range(10):
                rows.append(f"{a},{b},{2 * int(a) + int(b) + repeat * 0.03}")
    plan = {
        "interface_mode": "professional",
        "dependent_variables": ["y"], "fixed_factors": ["A", "B"],
        "method": "oneway_anova", "factor_combinations_enabled": True,
        "factor_combination_min_order": 1, "factor_combination_max_order": 1,
        "factor_combination_labels": {"A": "主处理模型", "B": "辅助处理模型"},
        "cross_model_p_adjust": "none",
    }
    client = TestClient(create_app(workspace_root=tmp_path / "workspace"))
    response = client.post(
        "/api/analyze",
        files={"file": ("none.csv", ("\n".join(rows) + "\n").encode(), "text/csv")},
        data={"plan_json": json.dumps(plan, ensure_ascii=False)},
    )
    assert response.status_code == 200, response.text
    payload = response.json()
    tests = [row for row in payload["result"]["summary"] if row["result_type"] == "test"]
    assert {row["factor_combination"] for row in tests} == {"主处理模型", "辅助处理模型"}
    assert {row["factor_columns"] for row in tests} == {"A", "B"}
    assert all(row["p_adjust_method"] == "none" for row in tests)
    assert all(row["p_adjusted_across_tasks"] is None for row in tests)
    assert all(row["significant_adjusted"] is None for row in tests)
    assert all(row["conclusion"].startswith("原始") for row in payload["result"]["overview"])
    assert all(row.get("task_id") is not None for row in payload["result"]["overview"])
    assert all(row["raw_conclusion"].startswith("原始") for row in payload["result"]["overview"])

    workbook_response = client.get(payload["batch_export"]["download_url"])
    workbook = openpyxl.load_workbook(io.BytesIO(workbook_response.content), read_only=True)
    headers = [cell.value for cell in next(workbook["结果总览"].iter_rows(min_row=3, max_row=3))]
    assert "组合名称" in headers
    assert "组合因素" in headers
    assert "校正前显著性" in headers
    assert "跨任务校正 p" not in headers
    assert "结论" not in headers
    assert "跨任务判断 p" not in headers


def test_order_parameter_expands_one_through_k_and_emits_standard_report(tmp_path):
    rows = ["A,B,C,D,y"]
    for a in ["0", "1"]:
        for b in ["0", "1"]:
            for c in ["0", "1"]:
                for d in ["0", "1"]:
                    for repeat in range(4):
                        rows.append(f"{a},{b},{c},{d},{int(a) + 2 * int(b) + 3 * int(c) + 4 * int(d) + repeat * 0.1}")
    plan = {
        "interface_mode": "professional",
        "dependent_variables": ["y"],
        "fixed_factors": ["A", "B", "C", "D"],
        "method": "multifactor_anova",
        "method_parameters": {"factor_model_order": 4},
        "factor_combinations_enabled": True,
        "factor_combination_order": 2,
    }
    client = TestClient(create_app(workspace_root=tmp_path / "workspace"))
    response = client.post(
        "/api/analyze",
        files={"file": ("order.csv", ("\n".join(rows) + "\n").encode(), "text/csv")},
        data={"plan_json": json.dumps(plan)},
    )
    assert response.status_code == 200, response.text
    payload = response.json()
    tasks = payload["result"]["results"]
    assert len(tasks) == 10  # C(4,1) + C(4,2)
    assert {item["subset_info"]["combination_order"] for item in tasks} == {1, 2}
    assert {item["result"]["method"]["name"] for item in tasks} == {"oneway_anova", "twoway_anova"}

    report = payload["batch_export"]["standardized_report"]
    assert report["schema_version"] == "datawork.standard-report.v1"
    archive_response = client.get(report["download_url"])
    assert archive_response.status_code == 200
    with zipfile.ZipFile(io.BytesIO(archive_response.content)) as archive:
        assert {"statistical_report.md", "canonical_result.json", "results.xlsx"} <= set(archive.namelist())
        canonical = json.loads(archive.read("canonical_result.json"))
    assert canonical["kind"] == "batch"
    assert canonical["schema_version"] == "datawork.standard-report.v1"
