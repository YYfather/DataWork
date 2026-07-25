"""批次分析工作簿导出。

工作簿按阅读顺序组织为结果总览、逐批结果、失败与警告、分析设置。
完整长表仍保留在运行结果中供机器处理，但不再作为默认 Excel 界面。
"""

from __future__ import annotations

import json
import unicodedata
from pathlib import Path
from typing import Any, Iterable
from uuid import uuid4

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from datawork.core.errors import DataWorkError, ErrorCode
from datawork.core.provenance import sha256_bytes
from datawork.engine.batch import BatchAnalysisResult, BatchResult
from datawork.engine.result import StatisticalResult


TITLE_FILL = PatternFill("solid", fgColor="173F5F")
HEADER_FILL = PatternFill("solid", fgColor="DCEAF4")
GOOD_FILL = PatternFill("solid", fgColor="E3F3E8")
NEUTRAL_FILL = PatternFill("solid", fgColor="F1F3F5")
ERROR_FILL = PatternFill("solid", fgColor="FBE4E6")

FIELD_LABELS = {
    "task_id": "任务",
    "factor_combination": "组合名称",
    "factor_columns": "组合因素",
    "combination_order": "组合阶数",
    "dependent_variable": "因变量",
    "dependent_variables": "联合因变量",
    "dependent_group": "联合因变量组",
    "dependent_combination": "因变量组合",
    "dependent_columns": "组合因变量",
    "dependent_combination_columns": "组合因变量（规范）",
    "dependent_combination_size": "因变量组合大小",
    "n": "样本量",
    "result_type": "结果类型",
    "effect": "效应/检验",
    "statistic_name": "统计量",
    "statistic_value": "统计量值",
    "f_value": "F 值",
    "approx_f": "近似 F",
    "df_num": "df1",
    "df_den": "df2",
    "sum_sq": "平方和 SS",
    "mean_sq": "均方 MS",
    "p_value": "原始 p",
    "p_adjusted": "校正 p",
    "p_adjusted_across_tasks": "跨任务校正 p",
    "significant": "原始显著",
    "significant_adjusted": "跨任务判断显著",
    "raw_conclusion": "校正前显著性",
    "conclusion": "结论",
    "status": "状态",
    "error": "失败原因",
    "estimate": "估计值",
    "se": "标准误",
    "ci_lower": "置信区间下限",
    "ci_upper": "置信区间上限",
    "eta_sq": "η²",
    "eta_sq_p": "偏 η²",
    "omega_sq": "ω²",
    "effect_size_name": "效应量",
    "effect_size_value": "效应量值",
    "is_significant": "显著",
    "term": "参数",
    "group": "水平/组",
    "mean": "均值/EMM",
    "letters": "字母",
    "method": "方法",
    "context": "条件",
    "detail": "说明",
    "test_name": "检查",
    "statistic": "统计量值",
    "passed": "状态",
    "name": "名称",
    "value": "数值",
    "fixed_factors": "分类因素",
    "random_factors": "随机分组",
    "random_slopes": "随机斜率",
    "covariates": "连续协变量",
    "subject_id": "对象 ID",
    "repeated_factor": "重复/时间因素",
    "split_by": "批量拆分列",
    "alpha": "显著性水平 α",
    "ss_type": "平方和类型",
    "equal_var": "假定方差齐性",
    "posthoc_method": "事后比较",
    "missing_policy": "缺失值策略",
    "research_goal": "研究目标",
    "test_value": "单样本参考值",
    "alternative": "备择假设",
    "expected_proportions": "理论比例",
    "method_parameters": "方法参数",
    "factor_combinations_enabled": "分类因素组合实验",
    "factor_combination_order": "计算阶数（最高阶）",
    "factor_combination_min_order": "最小组合阶数",
    "factor_combination_max_order": "最大组合阶数",
    "factor_combination_labels": "自定义组合名称",
    "dependent_task_mode": "联合响应任务模式",
    "dependent_variable_groups": "手工联合因变量组",
    "dependent_combination_min_size": "最小因变量组合大小",
    "dependent_combination_max_size": "最大因变量组合大小",
    "dependent_combination_labels": "自定义因变量组合名称",
    "combination_p_adjust": "跨任务判断方法",
    "cross_model_p_adjust": "跨模型判断方法",
    "estimate_marginal_means": "估计边际均值",
    "emm_factors": "EMM 因素",
    "contrast_correction": "对比校正",
    "diagnostic_plots": "生成诊断图",
    "interface_mode": "界面模式",
    "calibration_enabled": "启用校正",
    "calibration_method": "校正策略",
    "calibration_columns": "校正列",
    "calibration_baseline_column": "校正基准列",
    "calibration_baseline_value": "校正基准值",
}

SECTION_FIELDS = (
    ("核心检验", "primary_tests"),
    ("总体效应检验", "omnibus_tests"),
    ("模型参数", "coefficients"),
    ("单变量跟进", "follow_up_tests"),
    ("显著性字母分组", "significance_letters"),
    ("条件简单效应", "simple_effects"),
    ("估计边际均值", "estimated_marginal_means"),
    ("成对比较", "contrasts"),
    ("假设与数据诊断", "diagnostics"),
    ("模型拟合", "fit_statistics"),
    ("描述统计", "descriptive_stats"),
)


DEPENDENT_COMBINATION_AGGREGATES = (
    ("多元总体检验", "omnibus_tests"),
    ("单变量跟进", "follow_up_tests"),
    ("事后比较", "contrasts"),
)


def dependent_combination_aggregate_names(result: BatchAnalysisResult) -> list[str]:
    """Return aggregate sheets that have at least one row."""
    if result.settings.get("dependent_task_mode") != "combinations":
        return []
    names: list[str] = []
    for title, field_name in DEPENDENT_COMBINATION_AGGREGATES:
        if any(
            item.result is not None and bool(getattr(item.result, field_name))
            for item in result.results
        ):
            names.append(title)
    return names


def _display_width(value: str) -> int:
    """Estimate Excel display width while accounting for CJK full-width text."""
    return sum(2 if unicodedata.east_asian_width(character) in {"W", "F", "A"} else 1 for character in value)


class GenericBatchExportService:
    """将批次结果导出为一个可逐批阅读的工作簿。"""

    def export_xlsx(self, result: BatchAnalysisResult, path: str | Path) -> Path:
        if result.overview_df is None and result.summary_df is None:
            raise DataWorkError(ErrorCode.REPORT_FAILED, "批次分析没有可导出的结果")
        path = Path(path).expanduser().resolve()
        path.parent.mkdir(parents=True, exist_ok=True)
        # 报告目录可能已经很深；短临时文件名避免 Windows 传统路径长度限制。
        temporary = path.with_name(f".~dw-{uuid4().hex[:8]}.xlsx")

        workbook = Workbook()
        workbook.remove(workbook.active)
        self._write_overview(workbook, result)
        self._write_dependent_combination_aggregates(workbook, result)
        for index, item in enumerate(result.results, start=1):
            self._write_batch_sheet(workbook, index, item)
        self._write_issues(workbook, result.results)
        self._write_settings(workbook, result.settings)
        workbook.save(temporary)
        temporary.replace(path)
        return path

    def _write_dependent_combination_aggregates(
        self,
        workbook: Workbook,
        result: BatchAnalysisResult,
    ) -> None:
        active = set(dependent_combination_aggregate_names(result))
        for title, field_name in DEPENDENT_COMBINATION_AGGREGATES:
            if title not in active:
                continue
            rows: list[dict[str, Any]] = []
            for index, item in enumerate(result.results, start=1):
                if item.result is None:
                    continue
                metadata = {"task_id": index, **item.subset_info, "n": item.n_rows}
                for record in getattr(item.result, field_name):
                    rows.append({**metadata, **record.model_dump(mode="json")})
            worksheet = workbook.create_sheet(title)
            keys = list(dict.fromkeys(key for row in rows for key in row))
            worksheet.append([FIELD_LABELS.get(key, key) for key in keys])
            for row in rows:
                worksheet.append([self._cell_value(row.get(key)) for key in keys])
            self._style_table(worksheet, header_row=1)
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = (
                f"A1:{get_column_letter(worksheet.max_column)}{worksheet.max_row}"
            )

    def export_serialized(self, payload: dict[str, Any], path: str | Path) -> Path:
        """从工作区持久化的批次 JSON 重建分层工作簿。"""
        results = []
        for item in payload.get("results") or []:
            raw_result = item.get("result")
            results.append(BatchResult(
                subset_key=str(item.get("subset_key") or ""),
                subset_info=dict(item.get("subset_info") or {}),
                n_rows=int(item.get("n_rows") or 0),
                omnibus_tests=list(item.get("omnibus_tests") or []),
                records=list(item.get("records") or []),
                error=str(item.get("error") or ""),
                result=StatisticalResult.model_validate(raw_result) if raw_result else None,
            ))
        batch = BatchAnalysisResult(
            split_cols=list(payload.get("split_cols") or []),
            method=str(payload.get("method") or ""),
            dv_col=str(payload.get("dv_col") or ""),
            factor_cols=list(payload.get("factor_cols") or []),
            results=results,
            summary_df=pd.DataFrame(payload.get("summary") or []),
            overview_df=pd.DataFrame(payload.get("overview") or []),
            settings=dict(payload.get("settings") or {}),
        )
        return self.export_xlsx(batch, path)

    def _write_overview(self, workbook: Workbook, result: BatchAnalysisResult) -> None:
        worksheet = workbook.create_sheet("结果总览")
        source = result.overview_df if result.overview_df is not None else result.summary_df
        frame = source.copy() if source is not None else pd.DataFrame()
        adjustment_method = str(
            result.settings.get("cross_model_p_adjust")
            or result.settings.get("combination_p_adjust")
            or "holm"
        ).lower()
        hidden = {"result_type", "significant", "significant_adjusted"}
        if adjustment_method == "none":
            if "raw_conclusion" not in frame.columns and "significant" in frame.columns:
                frame["raw_conclusion"] = frame["significant"].map(
                    {True: "原始显著", False: "原始不显著"}
                )
            hidden.update({"p_adjusted_across_tasks", "conclusion"})
        frame = frame[[column for column in frame.columns if column not in hidden]]
        frame = frame.rename(columns={key: FIELD_LABELS.get(key, key) for key in frame.columns})
        worksheet.append(["批次分析结果总览"])
        worksheet["A1"].font = Font(bold=True, color="FFFFFF", size=14)
        worksheet["A1"].fill = TITLE_FILL
        worksheet.append([
            f"共 {len(result.results)} 个任务；未执行跨任务校正，结论按原始 p 判断。"
            if adjustment_method == "none"
            else f"共 {len(result.results)} 个任务；同时展示校正前判断，正式结论采用跨任务校正结果。"
        ])
        if frame.empty:
            worksheet.append(["暂无可展示结果"])
            return
        worksheet.append(list(frame.columns))
        for row in frame.itertuples(index=False, name=None):
            worksheet.append([self._cell_value(value) for value in row])
        if worksheet.max_column > 1:
            worksheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=worksheet.max_column)
        self._style_table(worksheet, header_row=3)
        worksheet.freeze_panes = "A4"
        worksheet.auto_filter.ref = f"A3:{get_column_letter(worksheet.max_column)}{worksheet.max_row}"
        conclusion_columns = [
            cell.column for cell in worksheet[3]
            if cell.value in {"校正前显著性", "结论"}
        ]
        for conclusion_column in conclusion_columns:
            letter = get_column_letter(conclusion_column)
            data_range = f"{letter}4:{letter}{worksheet.max_row}"
            worksheet.conditional_formatting.add(data_range, CellIsRule(operator="equal", formula=['"校正后显著"'], fill=GOOD_FILL))
            worksheet.conditional_formatting.add(data_range, CellIsRule(operator="equal", formula=['"原始显著"'], fill=GOOD_FILL))
            worksheet.conditional_formatting.add(data_range, CellIsRule(operator="equal", formula=['"执行失败"'], fill=ERROR_FILL))

    def _write_batch_sheet(self, workbook: Workbook, index: int, item: BatchResult) -> None:
        worksheet = workbook.create_sheet(f"批次_{index:03d}")
        worksheet.append([f"批次 {index:03d}"])
        worksheet["A1"].font = Font(bold=True, color="FFFFFF", size=14)
        worksheet["A1"].fill = TITLE_FILL
        worksheet.append(["任务标识", item.subset_key])
        worksheet.append(["样本量", item.n_rows])
        for key, value in item.subset_info.items():
            worksheet.append([FIELD_LABELS.get(key, key), self._cell_value(value)])
        if item.error:
            worksheet.append([])
            worksheet.append(["执行失败", item.error])
            worksheet.cell(worksheet.max_row, 1).fill = ERROR_FILL
            self._finish_sheet(worksheet)
            return
        if item.result is None:
            worksheet.append([])
            worksheet.append(["结果不可用"])
            self._finish_sheet(worksheet)
            return

        payload = item.result.model_dump(mode="json")
        for title, field_name in SECTION_FIELDS:
            rows = payload.get(field_name) or []
            if rows:
                self._append_section(worksheet, title, rows)
        warnings = payload.get("warnings") or []
        if warnings:
            self._append_section(worksheet, "分析提醒", [{"detail": warning} for warning in warnings])
        self._finish_sheet(worksheet)

    def _write_issues(self, workbook: Workbook, results: list[BatchResult]) -> None:
        worksheet = workbook.create_sheet("失败与警告")
        worksheet.append(["任务", "批次", "类型", "内容"])
        rows = 0
        for index, item in enumerate(results, start=1):
            if item.error:
                worksheet.append([index, item.subset_key, "失败", item.error])
                rows += 1
            if item.result is not None:
                for warning in item.result.warnings:
                    worksheet.append([index, item.subset_key, "警告", warning])
                    rows += 1
        if rows == 0:
            worksheet.append([None, None, "正常", "所有批次均未记录失败或警告"])
        self._style_table(worksheet, header_row=1)
        worksheet.freeze_panes = "A2"

    def _write_settings(self, workbook: Workbook, settings: dict[str, Any]) -> None:
        worksheet = workbook.create_sheet("分析设置")
        worksheet.append(["设置项", "取值"])
        for key, value in settings.items():
            worksheet.append([FIELD_LABELS.get(key, key), self._cell_value(value)])
        self._style_table(worksheet, header_row=1)
        worksheet.freeze_panes = "A2"

    def _append_section(self, worksheet, title: str, rows: Iterable[dict[str, Any]]) -> None:
        normalized = [dict(row) for row in rows]
        if not normalized:
            return
        worksheet.append([])
        worksheet.append([title])
        title_row = worksheet.max_row
        worksheet.cell(title_row, 1).font = Font(bold=True, color="FFFFFF")
        worksheet.cell(title_row, 1).fill = TITLE_FILL
        keys = list(dict.fromkeys(key for row in normalized for key in row.keys()))
        worksheet.append([FIELD_LABELS.get(key, key) for key in keys])
        header_row = worksheet.max_row
        for cell in worksheet[header_row]:
            cell.font = Font(bold=True, color="173F5F")
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for row in normalized:
            worksheet.append([self._cell_value(row.get(key)) for key in keys])

    def _style_table(self, worksheet, *, header_row: int) -> None:
        for cell in worksheet[header_row]:
            cell.font = Font(bold=True, color="173F5F")
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        self._finish_sheet(worksheet)

    @staticmethod
    def _finish_sheet(worksheet) -> None:
        worksheet.sheet_view.showGridLines = False
        for column_cells in worksheet.iter_cols():
            values = [str(cell.value) if cell.value is not None else "" for cell in column_cells[:200]]
            width = min(max(max((_display_width(value) for value in values), default=0) + 2, 10), 42)
            worksheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width
        for row in worksheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        if worksheet.max_column > 1:
            for row_number in range(1, worksheet.max_row + 1):
                first = worksheet.cell(row_number, 1)
                if (
                    first.value
                    and first.font.bold
                    and first.fill.fill_type == "solid"
                    and all(worksheet.cell(row_number, column).value is None for column in range(2, worksheet.max_column + 1))
                ):
                    worksheet.merge_cells(
                        start_row=row_number, start_column=1,
                        end_row=row_number, end_column=worksheet.max_column,
                    )

    @staticmethod
    def _cell_value(value: Any) -> Any:
        if value is None or (isinstance(value, float) and pd.isna(value)):
            return None
        if isinstance(value, bool):
            return "是" if value else "否"
        if isinstance(value, (list, tuple, dict)):
            return json.dumps(value, ensure_ascii=False, sort_keys=True)
        return value

    @staticmethod
    def file_payload(path: Path) -> dict[str, Any]:
        content = path.read_bytes()
        return {
            "filename": path.name,
            "size_bytes": len(content),
            "sha256": sha256_bytes(content),
        }
